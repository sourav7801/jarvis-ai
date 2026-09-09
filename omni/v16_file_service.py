"""Managed file ingestion for JARVIS V16.

The service owns a private runtime file area instead of granting agents arbitrary
filesystem access.  Uploaded bytes receive a stable file_id, SHA-256 identity,
metadata, workspace binding, and bounded text extraction for search.

No OCR is attempted here.  Binary document parsing uses installed local readers
when available and otherwise reports the parser as unavailable.
"""

from __future__ import annotations

from dataclasses import asdict, dataclass
from datetime import datetime, timezone
import csv
import hashlib
import io
import json
import mimetypes
import os
from pathlib import Path
import re
import shutil
import sqlite3
from threading import RLock
from typing import Any, Iterable
import uuid


PROJECT_ROOT = Path(__file__).resolve().parents[1]
DEFAULT_ROOT = PROJECT_ROOT / "data" / "workspaces" / "files"
DEFAULT_MAX_BYTES = 100 * 1024 * 1024
MAX_EXTRACTED_TEXT = 2_000_000

_ALLOWED_EXTENSIONS = {
    ".pdf", ".docx", ".xlsx", ".xls", ".csv", ".pptx",
    ".txt", ".md", ".json", ".xml", ".html", ".htm",
    ".jpg", ".jpeg", ".png", ".webp", ".gif",
    ".zip", ".py", ".js", ".ts", ".tsx", ".jsx", ".sql",
    ".ps1", ".bat", ".cmd", ".toml", ".yaml", ".yml", ".log",
}
_TEXT_EXTENSIONS = {
    ".txt", ".md", ".json", ".xml", ".html", ".htm", ".csv",
    ".py", ".js", ".ts", ".tsx", ".jsx", ".sql", ".ps1", ".bat",
    ".cmd", ".toml", ".yaml", ".yml", ".log",
}
_SAFE_WORKSPACE = re.compile(r"^[A-Z0-9][A-Z0-9_-]{0,63}$")


@dataclass(frozen=True)
class FileRecord:
    file_id: str
    original_name: str
    stored_name: str
    extension: str
    mime_type: str
    size_bytes: int
    sha256: str
    workspace: str
    created_at: str
    parser_status: str
    text_chars: int

    def to_dict(self) -> dict[str, Any]:
        return asdict(self)


class ManagedFileServiceV16:
    def __init__(self, root: Path | str = DEFAULT_ROOT, *, max_bytes: int = DEFAULT_MAX_BYTES) -> None:
        self.root = Path(root).resolve()
        self.max_bytes = max(1, int(max_bytes))
        self.objects = self.root / "objects"
        self.db_path = self.root / "index.sqlite3"
        self._lock = RLock()
        self.objects.mkdir(parents=True, exist_ok=True)
        self._ensure_db()

    @staticmethod
    def normalize_workspace(workspace: str | None) -> str:
        value = str(workspace or "HOME").strip().upper().replace(" ", "_")
        if not _SAFE_WORKSPACE.fullmatch(value):
            raise ValueError("workspace must contain only letters, numbers, '_' or '-'")
        return value

    @staticmethod
    def sanitize_filename(filename: str) -> str:
        raw = str(filename or "").strip()
        if not raw:
            raise ValueError("filename is required")
        # Never accept a caller-provided path.  Only its final component is a
        # legitimate uploaded filename, and traversal tokens are rejected.
        if raw != Path(raw).name or "/" in raw or "\\" in raw or raw in {".", ".."}:
            raise ValueError("filename must not contain a path")
        cleaned = re.sub(r"[^A-Za-z0-9._()\- ]+", "_", raw).strip(" .")
        if not cleaned:
            raise ValueError("filename is invalid after sanitization")
        if len(cleaned) > 180:
            stem = Path(cleaned).stem[:140]
            suffix = Path(cleaned).suffix[:20]
            cleaned = stem + suffix
        return cleaned

    @staticmethod
    def _extension(filename: str) -> str:
        return Path(filename).suffix.lower()

    def _validate_type(self, filename: str) -> str:
        extension = self._extension(filename)
        if extension not in _ALLOWED_EXTENSIONS:
            raise ValueError(f"unsupported upload type: {extension or '<none>'}")
        return extension

    def _connect(self) -> sqlite3.Connection:
        connection = sqlite3.connect(str(self.db_path), timeout=20.0)
        connection.row_factory = sqlite3.Row
        return connection

    def _ensure_db(self) -> None:
        self.root.mkdir(parents=True, exist_ok=True)
        connection = self._connect()
        try:
            connection.execute(
                """
                CREATE TABLE IF NOT EXISTS files (
                    file_id TEXT PRIMARY KEY,
                    original_name TEXT NOT NULL,
                    stored_name TEXT NOT NULL,
                    extension TEXT NOT NULL,
                    mime_type TEXT NOT NULL,
                    size_bytes INTEGER NOT NULL,
                    sha256 TEXT NOT NULL,
                    workspace TEXT NOT NULL,
                    created_at TEXT NOT NULL,
                    parser_status TEXT NOT NULL,
                    extracted_text TEXT NOT NULL DEFAULT '',
                    text_chars INTEGER NOT NULL DEFAULT 0,
                    UNIQUE(workspace, sha256, original_name)
                )
                """
            )
            connection.execute("CREATE INDEX IF NOT EXISTS idx_files_workspace_created ON files(workspace, created_at DESC)")
            connection.execute("CREATE INDEX IF NOT EXISTS idx_files_sha ON files(sha256)")
            connection.commit()
        finally:
            connection.close()

    @staticmethod
    def _sha256_bytes(content: bytes) -> str:
        return hashlib.sha256(content).hexdigest()

    def _object_path(self, file_id: str, stored_name: str) -> Path:
        directory = (self.objects / file_id[:2] / file_id).resolve()
        if self.objects not in directory.parents:
            raise RuntimeError("managed object path escaped file root")
        directory.mkdir(parents=True, exist_ok=True)
        return directory / stored_name

    @staticmethod
    def _row_to_record(row: sqlite3.Row) -> FileRecord:
        return FileRecord(
            file_id=str(row["file_id"]),
            original_name=str(row["original_name"]),
            stored_name=str(row["stored_name"]),
            extension=str(row["extension"]),
            mime_type=str(row["mime_type"]),
            size_bytes=int(row["size_bytes"]),
            sha256=str(row["sha256"]),
            workspace=str(row["workspace"]),
            created_at=str(row["created_at"]),
            parser_status=str(row["parser_status"]),
            text_chars=int(row["text_chars"]),
        )

    def _existing(self, workspace: str, digest: str, original_name: str) -> FileRecord | None:
        connection = self._connect()
        try:
            row = connection.execute(
                "SELECT * FROM files WHERE workspace=? AND sha256=? AND original_name=?",
                (workspace, digest, original_name),
            ).fetchone()
            return self._row_to_record(row) if row else None
        finally:
            connection.close()

    def ingest_bytes(
        self,
        filename: str,
        content: bytes | bytearray,
        *,
        workspace: str = "HOME",
    ) -> dict[str, Any]:
        safe_name = self.sanitize_filename(filename)
        extension = self._validate_type(safe_name)
        body = bytes(content)
        if not body:
            raise ValueError("uploaded file is empty")
        if len(body) > self.max_bytes:
            raise ValueError(f"file exceeds V16 upload limit of {self.max_bytes} bytes")
        workspace_name = self.normalize_workspace(workspace)
        digest = self._sha256_bytes(body)
        with self._lock:
            existing = self._existing(workspace_name, digest, safe_name)
            if existing is not None:
                return {"success": True, "duplicate": True, "file": existing.to_dict()}

            file_id = uuid.uuid4().hex
            stored_name = safe_name
            target = self._object_path(file_id, stored_name)
            temporary = target.with_name(f".{target.name}.{os.getpid()}.{uuid.uuid4().hex}.tmp")
            temporary.write_bytes(body)
            os.replace(temporary, target)

            extracted_text, parser_status = self._extract_text(target, extension)
            extracted_text = extracted_text[:MAX_EXTRACTED_TEXT]
            mime_type = mimetypes.guess_type(safe_name)[0] or "application/octet-stream"
            created = datetime.now(timezone.utc).isoformat()
            connection = self._connect()
            try:
                connection.execute(
                    """
                    INSERT INTO files (
                        file_id, original_name, stored_name, extension, mime_type,
                        size_bytes, sha256, workspace, created_at, parser_status,
                        extracted_text, text_chars
                    ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                    """,
                    (
                        file_id, safe_name, stored_name, extension, mime_type,
                        len(body), digest, workspace_name, created, parser_status,
                        extracted_text, len(extracted_text),
                    ),
                )
                connection.commit()
            except Exception:
                try:
                    shutil.rmtree(target.parent, ignore_errors=True)
                finally:
                    raise
            finally:
                connection.close()

        record = self.get(file_id)
        return {"success": True, "duplicate": False, "file": record["file"]}

    def register_path(
        self,
        source: Path | str,
        *,
        original_name: str | None = None,
        workspace: str = "HOME",
    ) -> dict[str, Any]:
        """Trusted-runtime helper for a file selected by the user.

        The file is copied into managed storage; agents subsequently receive only
        the resulting file_id, never broad access to the source directory.
        """

        path = Path(source).expanduser().resolve()
        if not path.is_file():
            raise FileNotFoundError(path)
        if path.stat().st_size > self.max_bytes:
            raise ValueError(f"file exceeds V16 upload limit of {self.max_bytes} bytes")
        return self.ingest_bytes(original_name or path.name, path.read_bytes(), workspace=workspace)

    def get(self, file_id: str) -> dict[str, Any]:
        identifier = str(file_id or "").strip()
        connection = self._connect()
        try:
            row = connection.execute("SELECT * FROM files WHERE file_id=?", (identifier,)).fetchone()
            if not row:
                return {"success": False, "reason": "FILE_NOT_FOUND", "file_id": identifier}
            record = self._row_to_record(row)
        finally:
            connection.close()
        path = self._object_path(record.file_id, record.stored_name)
        return {
            "success": True,
            "file": record.to_dict(),
            "managed_path": str(path),
            "exists": path.is_file(),
        }

    def read_bytes(self, file_id: str, *, max_bytes: int | None = None) -> bytes:
        result = self.get(file_id)
        if not result.get("success") or not result.get("exists"):
            raise FileNotFoundError(file_id)
        path = Path(str(result["managed_path"]))
        limit = self.max_bytes if max_bytes is None else min(max(1, int(max_bytes)), self.max_bytes)
        if path.stat().st_size > limit:
            raise ValueError("file exceeds requested read limit")
        return path.read_bytes()

    def extracted_text(self, file_id: str, *, max_chars: int = 200_000) -> dict[str, Any]:
        connection = self._connect()
        try:
            row = connection.execute(
                "SELECT original_name, parser_status, extracted_text FROM files WHERE file_id=?",
                (str(file_id),),
            ).fetchone()
            if not row:
                return {"success": False, "reason": "FILE_NOT_FOUND", "file_id": str(file_id)}
            text = str(row["extracted_text"] or "")[: max(1, int(max_chars))]
            return {
                "success": True,
                "file_id": str(file_id),
                "filename": str(row["original_name"]),
                "parser_status": str(row["parser_status"]),
                "text": text,
                "truncated": len(str(row["extracted_text"] or "")) > len(text),
            }
        finally:
            connection.close()

    def list_files(self, *, workspace: str | None = None, limit: int = 100) -> dict[str, Any]:
        cap = max(1, min(int(limit), 500))
        connection = self._connect()
        try:
            if workspace is None:
                rows = connection.execute("SELECT * FROM files ORDER BY created_at DESC LIMIT ?", (cap,)).fetchall()
            else:
                workspace_name = self.normalize_workspace(workspace)
                rows = connection.execute(
                    "SELECT * FROM files WHERE workspace=? ORDER BY created_at DESC LIMIT ?",
                    (workspace_name, cap),
                ).fetchall()
            records = [self._row_to_record(row).to_dict() for row in rows]
        finally:
            connection.close()
        return {"success": True, "count": len(records), "files": records}

    def search(self, query: str, *, workspace: str | None = None, limit: int = 50) -> dict[str, Any]:
        term = str(query or "").strip()
        if not term:
            return self.list_files(workspace=workspace, limit=limit)
        pattern = f"%{term.replace('%', '').replace('_', '')}%"
        cap = max(1, min(int(limit), 200))
        connection = self._connect()
        try:
            if workspace is None:
                rows = connection.execute(
                    "SELECT * FROM files WHERE original_name LIKE ? OR extracted_text LIKE ? ORDER BY created_at DESC LIMIT ?",
                    (pattern, pattern, cap),
                ).fetchall()
            else:
                workspace_name = self.normalize_workspace(workspace)
                rows = connection.execute(
                    "SELECT * FROM files WHERE workspace=? AND (original_name LIKE ? OR extracted_text LIKE ?) ORDER BY created_at DESC LIMIT ?",
                    (workspace_name, pattern, pattern, cap),
                ).fetchall()
            records = [self._row_to_record(row).to_dict() for row in rows]
        finally:
            connection.close()
        return {"success": True, "query": term, "count": len(records), "files": records}

    def _extract_text(self, path: Path, extension: str) -> tuple[str, str]:
        try:
            if extension in _TEXT_EXTENSIONS:
                if extension == ".csv":
                    return self._extract_csv(path), "PARSED"
                raw = path.read_text(encoding="utf-8", errors="replace")
                if extension == ".json":
                    try:
                        raw = json.dumps(json.loads(raw), ensure_ascii=False, indent=2)
                    except json.JSONDecodeError:
                        return raw, "PARSED_WITH_JSON_ERROR"
                return raw, "PARSED"
            if extension == ".xlsx":
                return self._extract_xlsx(path), "PARSED"
            if extension == ".docx":
                return self._extract_docx(path), "PARSED"
            if extension == ".pptx":
                return self._extract_pptx(path), "PARSED"
            if extension == ".pdf":
                return self._extract_pdf(path), "PARSED"
            if extension in {".jpg", ".jpeg", ".png", ".webp", ".gif"}:
                return "", "IMAGE_STORED_OCR_NOT_RUN"
            if extension in {".zip", ".xls"}:
                return "", "STORED_PARSER_NOT_ENABLED"
            return "", "STORED"
        except ImportError as exc:
            return "", f"PARSER_DEPENDENCY_UNAVAILABLE:{type(exc).__name__}"
        except Exception as exc:
            return "", f"PARSER_ERROR:{type(exc).__name__}:{str(exc)[:120]}"

    @staticmethod
    def _extract_csv(path: Path) -> str:
        pieces: list[str] = []
        with path.open("r", encoding="utf-8-sig", errors="replace", newline="") as handle:
            for index, row in enumerate(csv.reader(handle)):
                pieces.append("\t".join(str(value) for value in row))
                if index >= 20_000 or sum(map(len, pieces)) >= MAX_EXTRACTED_TEXT:
                    break
        return "\n".join(pieces)

    @staticmethod
    def _extract_xlsx(path: Path) -> str:
        from openpyxl import load_workbook

        workbook = load_workbook(path, read_only=True, data_only=False)
        try:
            chunks: list[str] = []
            for sheet in workbook.worksheets:
                chunks.append(f"[SHEET] {sheet.title}")
                for row in sheet.iter_rows(values_only=True):
                    chunks.append("\t".join("" if value is None else str(value) for value in row))
                    if sum(map(len, chunks)) >= MAX_EXTRACTED_TEXT:
                        return "\n".join(chunks)
            return "\n".join(chunks)
        finally:
            workbook.close()

    @staticmethod
    def _extract_docx(path: Path) -> str:
        from docx import Document

        document = Document(path)
        return "\n".join(paragraph.text for paragraph in document.paragraphs if paragraph.text)

    @staticmethod
    def _extract_pptx(path: Path) -> str:
        from pptx import Presentation

        presentation = Presentation(path)
        chunks: list[str] = []
        for index, slide in enumerate(presentation.slides, start=1):
            chunks.append(f"[SLIDE {index}]")
            for shape in slide.shapes:
                text = getattr(shape, "text", None)
                if text:
                    chunks.append(str(text))
        return "\n".join(chunks)

    @staticmethod
    def _extract_pdf(path: Path) -> str:
        try:
            from pypdf import PdfReader
        except ImportError:
            from PyPDF2 import PdfReader

        reader = PdfReader(str(path))
        chunks: list[str] = []
        for page in reader.pages:
            chunks.append(str(page.extract_text() or ""))
            if sum(map(len, chunks)) >= MAX_EXTRACTED_TEXT:
                break
        return "\n".join(chunks)

    def status(self) -> dict[str, Any]:
        return {
            "success": True,
            "version": "16.0",
            "service": "JARVIS_V16_MANAGED_FILE_SERVICE",
            "root": str(self.root),
            "allowed_extensions": sorted(_ALLOWED_EXTENSIONS),
            "max_upload_bytes": self.max_bytes,
            "managed_file_ids": True,
            "arbitrary_agent_filesystem_access": False,
            "ocr_automatic": False,
        }


MANAGED_FILE_SERVICE_V16 = ManagedFileServiceV16()


__all__ = ["FileRecord", "ManagedFileServiceV16", "MANAGED_FILE_SERVICE_V16"]
